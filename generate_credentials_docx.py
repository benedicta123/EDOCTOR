import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PRIMARY_GREEN = RGBColor(82, 153, 39)     # #529927
DARK_NAVY = RGBColor(19, 42, 69)          # #132A45
TEXT_MUTED = RGBColor(100, 116, 139)      # #64748B
DARK_GRAY = RGBColor(30, 41, 59)          # #1E293B

def set_cell_background(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_credentials_docx(docx_path):
    doc = Document()

    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

        # Header
        hp = section.header.paragraphs[0]
        hp.text = "eDoctor (Togo Health Direct) — Guide des Identifiants de Test"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if hp.runs:
            hp.runs[0].font.size = Pt(8.5)
            hp.runs[0].font.color.rgb = TEXT_MUTED

        # Footer
        fp = section.footer.paragraphs[0]
        fp.text = "Document Confidentiel de Recette • Plateforme eDoctor 2026"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if fp.runs:
            fp.runs[0].font.size = Pt(8.5)
            fp.runs[0].font.color.rgb = TEXT_MUTED

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("🔑 eDoctor — Annuaire des Comptes & Identifiants de Test")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_GREEN

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    run_sub = p_sub.add_run("Base de données PostgreSQL : 'edoctor' • Tous les comptes existants classés par rôle et interface")
    run_sub.font.size = Pt(10.5)
    run_sub.font.color.rgb = TEXT_MUTED

    # Info Callout Box (Password commun & Sécurité 2FA)
    table_info = doc.add_table(rows=1, cols=1)
    table_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_info = table_info.rows[0].cells[0]
    set_cell_background(c_info, "F0FDF4") # soft green
    set_cell_margins(c_info, top=120, bottom=120, left=180, right=180)
    p_info = c_info.paragraphs[0]
    p_info.paragraph_format.space_after = Pt(2)
    r1 = p_info.add_run("📌 MOT DE PASSE COMMUN POUR LES COMPTES DE RECETTE :  ")
    r1.font.bold = True
    r1.font.color.rgb = DARK_NAVY
    r2 = p_info.add_run("Password123!")
    r2.font.bold = True
    r2.font.name = 'Consolas'
    r2.font.size = Pt(12)
    r2.font.color.rgb = PRIMARY_GREEN
    
    p_secu = c_info.add_paragraph()
    p_secu.paragraph_format.space_before = Pt(4)
    p_secu.paragraph_format.space_after = Pt(2)
    r_secu = p_secu.add_run("🔐 SÉCURITÉ RENFORCÉE & DOUBLE AUTHENTIFICATION (2FA OTP) :\n")
    r_secu.font.bold = True
    r_secu.font.size = Pt(9.5)
    r_secu.font.color.rgb = DARK_NAVY
    r_secu_desc = p_secu.add_run(
        "• Double Authentification (2FA) : À la connexion, un code OTP à 6 chiffres est instantanément expédié à l'adresse email du compte via la passerelle transactionnelle Brevo (contacts@ewaregroup.org). En mode développement local (APP_DEBUG=true), ce code est également retourné dans la réponse JSON pour accélérer les tests.\n"
        "• Protection anti-brute force : Après 10 échecs de mot de passe consécutifs, le compte est automatiquement verrouillé pendant 30 minutes et un email d'alerte de sécurité est transmis.\n"
        "• Complexité requise : 8 caractères minimum comprenant obligatoirement majuscule, minuscule, chiffre et caractère spécial (ex: Password123! ou Ed0ctor2026!#UltraSecure)."
    )
    r_secu_desc.font.size = Pt(8.5)
    r_secu_desc.font.color.rgb = DARK_GRAY

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(6)

    # SECTIONS DATA
    sections_data = [
        {
            "role_title": "🩺 1. MÉDECINS PRATICIENS (Portail Praticien)",
            "app_target": "Application : edoctor_praticient • Interface : Téléconsultation vidéo, Dossier Patient, Prescriptions Ordonnances & Analyses Labo",
            "headers": ["Nom du Praticien", "Spécialité & Hôpital", "Email de Connexion", "Mot de Passe", "Téléphone"],
            "rows": [
                ["Dr. Paul Koffi", "Médecine Générale\nCHU de Démo", "dr.koffi@edoctor.test", "Password123!", "+225 07 00 00 02 01"],
                ["Dr. Amina Touré", "Pédiatrie\nCHU de Démo", "dr.amina@edoctor.test", "Password123!", "+225 07 00 00 02 02"],
            ]
        },
        {
            "role_title": "💊 2. PHARMACIENS TITULAIRES & OFFICINES (Portail Pharmacie)",
            "app_target": "Application : edoctor_pharmacy • Interface : Commandes en ligne, Scan QR Code Ordonnance, Dispensation, Import/Export Stocks Excel & CSV",
            "headers": ["Pharmacien Titulaire", "Pharmacie Rattachée", "Email de Connexion", "Mot de Passe", "Adresse / Ville"],
            "rows": [
                ["Dr. Koffi Mensah", "Grande Pharmacie du Golfe\n(ID Officine: 6)", "pharmacien.test@edoctor.tg", "Password123!", "Bd du 13 Janvier, Dékcon\nLomé, Togo"],
                ["Ange AMEGAH", "PHARMACIE AVEPOZO\n(ID Officine: 4)", "hounkanliemeraud@gmail.com", "Password123!", "Avépozo\nLomé, Togo"],
                ["Dr. Luc Bernard", "Grande Pharmacie Centrale\n(ID Officine: 1)", "pharmacien.luc@edoctor.test", "Password123!", "Avenue Chardy, Plateau\nAbidjan"],
                ["Dr. Fatou Diallo", "Pharmacie de la Paix\n(ID Officine: 2)", "pharmacienne.fatou@edoctor.test", "Password123!", "Bd de Marseille, Zone 4\nAbidjan"],
                ["Dr. Marc Yao", "Pharmacie Sainte-Marie\n(ID Officine: 3)", "pharmacien.marc@edoctor.test", "Password123!", "Cocody Angré\nAbidjan"]
            ]
        },
        {
            "role_title": "👤 3. PATIENTS (Application Mobile eDoctor)",
            "app_target": "Application : edoctor_mobile • Interface : Téléconsultation avec praticiens disponibles, Ordonnances certifiées PDF, Bons d'analyses Labo, Commandes Officines, Soins à Domicile",
            "headers": ["Nom Complet", "Profil Patient", "Email de Connexion", "Mot de Passe", "Téléphone"],
            "rows": [
                ["Bénédicta Hounkanli", "Compte Patient Principal\nAntécédents / Carnet vérifié", "benedictehounkanli@gmail.com", "Password123!", "+228 90 00 00 00"],
                ["Jean Dupont", "Patient Démo 1\nHistorique consultations", "patient.jean@edoctor.test", "Password123!", "+225 07 00 00 03 01"],
                ["Marie Kouassi", "Patiente Démo 2\nSuivi pédiatrique", "patiente.marie@edoctor.test", "Password123!", "+225 07 00 00 03 02"]
            ]
        },
        {
            "role_title": "🩹 4. PERSONNEL INFIRMIER (Soins à Domicile)",
            "app_target": "Application : edoctor_praticient • Interface : Tournées soignantes, Prélèvements, Pansements, Validations à domicile",
            "headers": ["Nom Infirmier(e)", "Établissement", "Email de Connexion", "Mot de Passe", "Téléphone"],
            "rows": [
                ["Infirmier Paul Yao", "CHU de Démo\nSoins à domicile", "infirmier.paul@edoctor.test", "Password123!", "+225 07 00 00 04 01"],
                ["Infirmière Sarah Koné", "CHU de Démo\nSoins & Prélèvements", "infirmiere.sarah@edoctor.test", "Password123!", "+225 07 00 00 04 02"]
            ]
        },
        {
            "role_title": "🏛️ 5. ADMINISTRATION HOSPITALIÈRE (CHU & Cliniques)",
            "app_target": "Application : edoctor_praticient (Menu Hôpital) • Interface : Gestion des équipes soignantes, Suivi des consultations, Coordination des soins",
            "headers": ["Nom Administrateur", "Rôle & Structure", "Email de Connexion", "Mot de Passe", "Téléphone"],
            "rows": [
                ["Dr. Kouamé Directeur", "Directeur Hospitalier\nCHU de Démo (ID: 1)", "admin.chu@edoctor.test", "Password123!", "+225 07 00 00 01 01"],
            ]
        },
        {
            "role_title": "👑 6. SUPER-ADMINISTRATION & GOUVERNANCE CENTRALE (Portail Dédié)",
            "app_target": "Application : edoctor_admin • Accès Web : https://admin-edoctor.ewaregroup.org (Docker Nginx SPA) • Interface : Agrément Hôpitaux & Pharmacies, Audit des Transactions, Litiges & Modération",
            "headers": ["Administrateur Système", "Rôle & Privilèges", "Email de Connexion", "Mot de Passe", "Téléphone"],
            "rows": [
                ["Dr. Kouamé / Équipe eDoctor", "Super-Administrateur National\nGouvernance Centrale & Agréments", "hounkanliemeraude@gmail.com", "Password123!", "+228 90 00 00 01"],
                ["Admin Technique eDoctor", "Super-Administrateur Système\nAudits, Sécurité & Dokploy", "admin@edoctor.test", "Password123!", "+228 91 00 00 00"]
            ]
        }
    ]

    for sec in sections_data:
        # Section Header
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(14)
        p_sec.paragraph_format.space_after = Pt(2)
        p_sec.paragraph_format.keep_with_next = True
        r_sec = p_sec.add_run(sec["role_title"])
        r_sec.font.name = 'Arial'
        r_sec.font.size = Pt(13)
        r_sec.font.bold = True
        r_sec.font.color.rgb = DARK_NAVY

        # Target App
        p_target = doc.add_paragraph()
        p_target.paragraph_format.space_after = Pt(6)
        p_target.paragraph_format.keep_with_next = True
        r_t = p_target.add_run(sec["app_target"])
        r_t.font.size = Pt(9.5)
        r_t.font.italic = True
        r_t.font.color.rgb = PRIMARY_GREEN

        # Table
        headers = sec["headers"]
        rows = sec["rows"]
        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        # Header Row
        hdr_row = table.rows[0]
        for c_idx, h_text in enumerate(headers):
            cell = hdr_row.cells[c_idx]
            cell.text = h_text
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            set_cell_background(cell, "132A45") # dark navy
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(255, 255, 255)

        # Data Rows
        for r_idx, r_data in enumerate(rows, 1):
            row = table.rows[r_idx]
            bg = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
            for c_idx, val in enumerate(r_data):
                cell = row.cells[c_idx]
                cell.text = val
                set_cell_background(cell, bg)
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.15
                if p.runs:
                    p.runs[0].font.size = Pt(9.0)
                    if c_idx == 2 or c_idx == 3: # Email / Password
                        p.runs[0].font.name = 'Consolas'
                        p.runs[0].font.bold = (c_idx == 3)
                        if c_idx == 3:
                            p.runs[0].font.color.rgb = PRIMARY_GREEN
                        else:
                            p.runs[0].font.color.rgb = DARK_GRAY
                    else:
                        p.runs[0].font.color.rgb = DARK_GRAY

        p_end = doc.add_paragraph()
        p_end.paragraph_format.space_after = Pt(4)

    doc.save(docx_path)
    print(f"Credentials document generated: {docx_path}")

if __name__ == "__main__":
    generate_credentials_docx(r"c:\Users\belou\Documents\edoctor\IDENTIFIANTS_UTILISATEURS_TESTS_EDOCTOR.docx")
    generate_credentials_docx(r"c:\Users\belou\Documents\edoctor\docs\IDENTIFIANTS_UTILISATEURS_TESTS_EDOCTOR.docx")
