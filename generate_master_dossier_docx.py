import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PRIMARY_GREEN = RGBColor(82, 153, 39)     # #529927 (Couleur Santé eDoctor)
DARK_NAVY = RGBColor(19, 42, 69)          # #132A45
TEXT_MUTED = RGBColor(100, 116, 139)      # #64748B
DARK_GRAY = RGBColor(30, 41, 59)          # #1E293B
CRIMSON_RED = RGBColor(220, 38, 38)       # #DC2626
LIGHT_BG = "F8FAFC"

def set_cell_background(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text, title="NOTE STRATÉGIQUE"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.8)
    cell = table.cell(0, 0)
    set_cell_background(cell, "ECFDF5") # vert très clair
    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
    
    cp = cell.paragraphs[0]
    r_t = cp.add_run(f"📌 {title} : ")
    r_t.font.name = 'Arial'
    r_t.font.bold = True
    r_t.font.size = Pt(10)
    r_t.font.color.rgb = PRIMARY_GREEN
    
    r_txt = cp.add_run(text)
    r_txt.font.name = 'Arial'
    r_txt.font.size = Pt(9.5)
    r_txt.font.color.rgb = DARK_GRAY

def create_master_dossier_docx(output_path):
    doc = Document()

    # Configuration des marges
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

        # En-tête
        hp = section.header.paragraphs[0]
        hp.text = "eDoctor (Togo Health Direct) — Dossier Stratégique & Procédures Opérationnelles 2026"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if hp.runs:
            hp.runs[0].font.size = Pt(8.5)
            hp.runs[0].font.color.rgb = TEXT_MUTED

        # Pied de page
        fp = section.footer.paragraphs[0]
        fp.text = "Document Confidentiel de Direction Générale • Plateforme eDoctor Togo"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if fp.runs:
            fp.runs[0].font.size = Pt(8.5)
            fp.runs[0].font.color.rgb = TEXT_MUTED

    # Page de Garde / Grand Titre
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(16)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("🏥 DOSSIER STRATÉGIQUE & PROCÉDURES OPÉRATIONNELLES")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = DARK_NAVY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(16)
    run_sub = p_sub.add_run("Monétisation Tripartite, Sécurisation des Flux Financiers, Logistique de Livraison & Verrous Anti-Contournement")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = PRIMARY_GREEN

    # Métadonnées du document
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(14)
    r_meta = p_meta.add_run(
        "Auteur : Direction Technique & Stratégique eDoctor  |  Destinataire : Direction Générale & Partenaires Hospitaliers  |  Date : Septembre 2026"
    )
    r_meta.font.size = Pt(9)
    r_meta.font.color.rgb = TEXT_MUTED

    add_callout(
        doc,
        "Ce dossier de synthèse rassemble l'ensemble des règles opérationnelles, économiques et juridiques de l'écosystème eDoctor. "
        "Il détaille précisément comment les Hôpitaux, les Pharmacies et la Startup génèrent leurs revenus, comment les flux Mobile Money "
        "sont séquestrés puis partagés, comment la logistique de livraison est orchestrée de bout en bout, et comment tout contournement est neutralisé.",
        "RÉSUMÉ EXÉCUTIF"
    )

    # ─────────────────────────────────────────────────────────────
    # SOMMAIRE
    # ─────────────────────────────────────────────────────────────
    p_toc = doc.add_paragraph()
    p_toc.paragraph_format.space_before = Pt(10)
    p_toc.paragraph_format.space_after = Pt(6)
    r_toc = p_toc.add_run("Table des Matières")
    r_toc.font.name = 'Arial'
    r_toc.font.size = Pt(14)
    r_toc.font.bold = True
    r_toc.font.color.rgb = DARK_NAVY

    toc_items = [
        "1. Modèle Économique Tripartite (Hôpitaux, Pharmacies, Startup)",
        "2. Rémunération des Hôpitaux & Cliniques (0 FCFA d'entrée, 85% de reversement)",
        "3. Revenus & Rentabilité des Pharmacies d'Officine (Conversion 95%, zéro impayé)",
        "4. Modèle de Profit de notre Startup eDoctor (Grille de commissions et simulations FCFA)",
        "5. Fonctionnement Bancaire : Agrégateur, Compte Séquestre (Escrow) et Split Payment",
        "6. Verrous Anti-Contournement & Anti-Fraude (Anti-WhatsApp, QR Code infalsifiable)",
        "7. Sécurisation des Soins à Domicile (Anti-cash, Code OTP, responsabilité hospitalière)",
        "8. Procédure Logistique de Livraison de Médicaments (Tarif GPS, coursiers, respect du secret médical)",
        "9. Gouvernance Centrale, Portail Super-Admin (edoctor_admin) & Sécurité FinTech (2FA Brevo, HMAC Webhook)",
        "10. Matrice Récapitulative des Risques & Garanties Opérationnelles"
    ]
    for item in toc_items:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(item)
        r.font.name = 'Arial'
        r.font.size = Pt(10)
        r.font.color.rgb = DARK_GRAY

    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────
    # CHAPITRE 1 : MODÈLE ÉCONOMIQUE TRIPARTITE
    # ─────────────────────────────────────────────────────────────
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    r = h.add_run("1. Modèle Économique Tripartite : Une Stratégie « Win-Win-Win »")
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = DARK_NAVY

    p = doc.add_paragraph()
    p.add_run(
        "Pour s'imposer durablement au Togo et dans la sous-région, eDoctor refuse tout modèle prédateur. "
        "La plateforme crée de la valeur économique réelle pour chaque partie prenante :\n\n"
        "• Pour le Patient : Accès immédiat 24/7 à un médecin agréé sans transport, devis transparent avant paiement, "
        "zéro monnaie à chercher à la livraison, dossier médical numérique universel.\n"
        "• Pour l'Hôpital / Clinique : Nouveaux revenus sur les heures creuses des médecins, valorisation du personnel soignant, "
        "flux continu de patients vers le plateau technique pour examens lourds, 0 FCFA de risque financier.\n"
        "• Pour la Pharmacie d'Officine : Encaissement immédiat garanti à 100% sans crédit client, extension de la zone de chalandise "
        "à toute l'agglomération grâce à la livraison, réduction des pertes sur médicaments périmés.\n"
        "• Pour notre Startup eDoctor : Marge nette prélevée automatiquement à la source sur chaque transaction (téléconsultation, "
        "médicaments, logistique, soins infirmiers) sans coût de stockage ni flotte lourde en propre."
    )

    # ─────────────────────────────────────────────────────────────
    # CHAPITRE 2 : RÈGLES FINANCIÈRES HÔPITAUX
    # ─────────────────────────────────────────────────────────────
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    r = h.add_run("2. Rémunération des Hôpitaux & Cliniques Partenaires")
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = DARK_NAVY

    p = doc.add_paragraph()
    p.add_run(
        "Règle n°1 : Adhésion à 0 FCFA d'investissement initial\n"
        "L'hôpital ou la clinique privée ne paie aucun droit d'entrée ni aucun abonnement obligatoire au démarrage. "
        "Le modèle est 100% au succès (Success-Fee) : eDoctor ne prélève sa commission que lorsque la structure encaisse réellement de l'argent.\n\n"
        "Règle n°2 : Liberté tarifaire & Reversement de 80% à 85%\n"
        "L'hôpital configure ses tarifs selon la spécialité (ex: 5 000 FCFA pour un médecin généraliste, 10 000 FCFA pour un pédiatre ou cardiologue). "
        "L'hôpital perçoit 80% à 85% du montant net directement dans son solde disponible dès la fin de l'acte.\n\n"
        "Règle n°3 : L'offre « eDoctor Pro » (Optionnelle pour plus tard)\n"
        "Seules les polycliniques souhaitant des fonctionnalités logicielles sur-mesure (interfaçage ERP, gestion des gardes de 50 soignants, "
        "sous-domaine web dédié) souscriront ultérieurement à un abonnement SaaS de 25 000 à 75 000 FCFA/mois. Ce n'est en aucun cas un prérequis au lancement."
    )

    # ─────────────────────────────────────────────────────────────
    # CHAPITRE 3 : RÈGLES FINANCIÈRES PHARMACIES
    # ─────────────────────────────────────────────────────────────
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    r = h.add_run("3. Revenus & Rentabilité des Pharmacies d'Officine")
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = DARK_NAVY

    p = doc.add_paragraph()
    p.add_run(
        "Règle n°1 : Taux de conversion de l'ordonnance proche de 100%\n"
        "Dans le circuit médical classique, près de 35% des ordonnances papier ne sont jamais achetées par le patient. "
        "Sur eDoctor, la consultation débouche directement sur le bouton « Commander mes médicaments » : l'ordonnance est solvabilisée sur-le-champ.\n\n"
        "Règle n°2 : Trésorerie 100% sécurisée sans impayé\n"
        "La pharmacie ne commence la préparation du sachet que lorsque le paiement total du panier est confirmé par T-Money/Flooz. "
        "L'officine touche 95% à 97% du montant brut de la vente sans aucun risque d'impayé.\n\n"
        "Règle n°3 : Respect strict du Code de Déontologie Pharmaceutique\n"
        "La commission eDoctor (3% à 5%) est strictement rémunérée au titre des frais technologiques d'intermédiation numérique et de mise en relation, "
        "conformément aux recommandations de l'Ordre National des Pharmaciens du Togo."
    )

    # ─────────────────────────────────────────────────────────────
    # CHAPITRE 4 : PROFIT DE NOTRE STARTUP (SIMULATIONS CHIFFRÉES)
    # ─────────────────────────────────────────────────────────────
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    r = h.add_run("4. Modèle de Profit de la Startup eDoctor & Grille Tarifaire")
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = DARK_NAVY

    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["Acte / Service", "Prix Facturé Patient", "Part Hôpital / Praticien", "Part Pharmacie / Coursier", "Marge Nette eDoctor"]
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "132A45")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = 'Arial'
            run.font.size = Pt(9)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    data_rows = [
        ("Téléconsultation Généraliste", "5 000 FCFA", "4 250 FCFA (85%)", "—", "750 FCFA (15%)"),
        ("Téléconsultation Spécialiste", "10 000 FCFA", "8 500 FCFA (85%)", "—", "1 500 FCFA (15%)"),
        ("Panier Pharmacie (Retrait)", "15 000 FCFA", "—", "14 400 FCFA (96%)", "600 FCFA (4%)"),
        ("Panier Pharmacie + Livraison", "16 200 FCFA", "—", "14 400 F (Pharm.) + 900 F (Coursier)", "900 FCFA (Com. + Marge livr.)"),
        ("Soin Infirmier à Domicile", "8 000 FCFA", "6 800 FCFA (85%)", "—", "1 200 FCFA (15%)")
    ]

    for row_idx, row_data in enumerate(data_rows):
        row = table.add_row()
        bg_color = LIGHT_BG if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            cell = row.cells[col_idx]
            cell.text = cell_value
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
            p = cell.paragraphs[0]
            if col_idx in [1, 2, 3, 4]:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if col_idx == 4:
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = PRIMARY_GREEN
            elif col_idx == 0:
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = DARK_NAVY
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(8.5)

    # ─────────────────────────────────────────────────────────────
    # CHAPITRE 5 : FONCTIONNEMENT BANCAIRE & SÉQUESTRE ESCROW
    # ─────────────────────────────────────────────────────────────
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    r = h.add_run("5. Procédure d'Encaissement, Séquestre (Escrow) et Split Payment")
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = DARK_NAVY

    p = doc.add_paragraph()
    p.add_run(
        "Étape 1 : Le Contrat Unique avec l'Agrégateur de Paiement (FedaPay, CinetPay ou PayDunya)\n"
        "Notre startup ouvre un Compte Marchand Entreprise (*Merchant Account*) auprès de l'agrégateur agréé. "
        "Ce compte unique fournit à la fois l'API d'encaissement Mobile Money (T-Money, Flooz) et le portefeuille central séquestre.\n\n"
        "Étape 2 : Le Paiement Préalable Obligatoire par le Patient (Timing T-0)\n"
        "Le patient est obligatoirement débité AVANT le démarrage de la consultation vidéo. Le médecin ne décroche jamais pour un appel non payé.\n\n"
        "Étape 3 : Le Compte Séquestre Temporaire (Escrow)\n"
        "L'argent encaissé est bloqué sur le compte séquestre eDoctor. Si la consultation est annulée ou si la connexion internet est défaillante "
        "avant la mise en relation, le patient est remboursé à 100% de manière automatisée.\n\n"
        "Étape 4 : Le Déclencheur d'Acte Légale (/end) et le Split Payment Instantané\n"
        "Dès que le médecin clique sur « Clôturer la consultation » et valide son diagnostic, l'API eDoctor découpe instantanément la somme : "
        "85% vont au portefeuille virtuel du praticien, 15% restent au profit d'eDoctor.\n\n"
        "Étape 5 : Le Déboursement / Retrait (Payout)\n"
        "L'hôpital ou le praticien déclenche le retrait de son solde vers son compte T-Money/Flooz ou compte bancaire directement depuis son tableau de bord."
    )

    # ─────────────────────────────────────────────────────────────
    # CHAPITRE 6 : VERROUS ANTI-CONTOURNEMENT (MÉDECINS)
    # ─────────────────────────────────────────────────────────────
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    r = h.add_run("6. Verrous Anti-Contournement : Comment Empêcher la Fraude Médicale (WhatsApp)")
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = DARK_NAVY

    p = doc.add_paragraph()
    p.add_run(
        "Pour éviter que les médecins et patients ne s'échangent leurs coordonnées pour continuer en direct hors plateforme, "
        "eDoctor intègre 5 verrous systématiques :\n\n"
        "1. L'Ordonnance Numérique Officielle avec QR Code Infalsifiable :\n"
        "Un médecin sur WhatsApp ne peut pas délivrer d'ordonnance reconnue par les pharmacies. Les officines partenaires exigent le code "
        "de référence de consultation eDoctor (ex: CNS-2609-CHUD-4DKE) et le QR Code officiel. Sans passage par la plateforme, pas d'ordonnance valide.\n\n"
        "2. Le Dossier Médical et la Prise en charge Assurance (INAM, Gras Savoye, Sunu) :\n"
        "Pour obtenir le remboursement employeur ou mutuelle, la feuille de soins électronique eDoctor est indispensable. "
        "Les échanges informels WhatsApp sont rejetés d'office par les compagnies d'assurance.\n\n"
        "3. Anonymisation Complète des Numéros Personnels :\n"
        "Durant la consultation, l'appel vidéo et le chat transitent via nos serveurs chiffrés (WebRTC). Le patient et le médecin ne voient jamais "
        "le numéro de téléphone de l'autre.\n\n"
        "4. La Garantie Absolue de Paiement pour le Praticien :\n"
        "En direct sur WhatsApp, les patients ne paient que dans 20% des cas. Sur eDoctor, le médecin a la garantie d'être payé à 100% car l'argent "
        "est déjà débité avant l'appel.\n\n"
        "5. Conventionnement Hospitalier & Sanctions Déontologiques :\n"
        "Les praticiens exercent sous la responsabilité de leur Hôpital de rattachement. Tout détournement avéré expose le praticien à une suspension "
        "immédiate et à un signalement auprès de l'Ordre National des Médecins."
    )

    # ─────────────────────────────────────────────────────────────
    # CHAPITRE 7 : SÉCURISATION SOINS À DOMICILE (ANTI-CASH)
    # ─────────────────────────────────────────────────────────────
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    r = h.add_run("7. Sécurisation des Soins à Domicile : Comment Empêcher le Cash en Main Propre")
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = DARK_NAVY

    p = doc.add_paragraph()
    p.add_run(
        "1. La Prestation est 100% prépayée sur l'application mobile :\n"
        "L'infirmier ne reçoit la fiche de mission avec l'adresse du patient que si le paiement (ex: 8 000 FCFA) a déjà été validé par T-Money. "
        "Le patient ayant déjà payé par mobile refusera formellement de verser des espèces à l'infirmier.\n\n"
        "2. Le Code Secret OTP du Patient :\n"
        "À l'arrivée de l'infirmier au domicile, le patient lui communique un code OTP à 4 chiffres généré sur son téléphone. "
        "C'est la saisie de ce code par l'infirmier qui prouve la visite et déclenche le paiement de ses honoraires.\n\n"
        "3. Rattachement Hospitalier de l'Infirmier :\n"
        "L'infirmier est rattaché à son hôpital conventionné. En cas de tentative d'extorsion de cash, le patient le signale dans l'onglet "
        "« Assistance & Réclamations », entraînant l'ouverture d'une enquête disciplinaire immédiate."
    )

    # ─────────────────────────────────────────────────────────────
    # CHAPITRE 8 : LOGISTIQUE DE LIVRAISON DE MÉDICAMENTS
    # ─────────────────────────────────────────────────────────────
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    r = h.add_run("8. Procédure Logistique de Livraison de Médicaments à Lomé")
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = DARK_NAVY

    p = doc.add_paragraph()
    p.add_run(
        "Règle n°1 : Tarification Algorithmique GPS Fixée par eDoctor (Jamais par le livreur)\n"
        "Le tarif est calculé automatiquement en fonction de la distance réelle entre l'officine et le patient :\n"
        "• Zone 1 (0 à 3 km) : 1 000 FCFA\n"
        "• Zone 2 (3 à 7 km, urbain Lomé) : 1 200 à 1 500 FCFA\n"
        "• Zone 3 (> 7 km, périphérie Grand Lomé) : 2 000 FCFA\n"
        "Le montant de la livraison est affiché en toute clarté dans le panier avant validation. Le patient paie tout en un seul débit Mobile Money.\n\n"
        "Règle n°2 : Modèle Logistique Hybride (Clé en main eDoctor vs Coursier Officine)\n"
        "• 90% des pharmacies n'ayant pas de coursiers, eDoctor gère la mise à disposition automatique d'un livreur partenaire dès que le paquet est prêt.\n"
        "• Les rares officines équipées de motos peuvent cocher l'option « Livré par nos soins » et conserver les frais de livraison.\n\n"
        "Règle n°3 : Recrutement & Animation des Coursiers de Santé eDoctor\n"
        "• Recrutement sur dossier : Permis A valide, casier judiciaire vierge, moto assurée, smartphone 4G.\n"
        "• Rémunération attractive : Sur 1 200 FCFA facturés, le coursier touche 900 FCFA nets instantanément par T-Money à la remise du colis.\n"
        "• Équipement sanitaire : Sac isotherme étanche pour protéger les médicaments contre la chaleur (chaîne du froid pour l'insuline).\n"
        "• Confidentialité totale : Le paquet est scellé et opaque. Le livreur ignore la nature des médicaments transportés (Secret Médical absolu)."
    )

    # ─────────────────────────────────────────────────────────────
    # CHAPITRE 9 : GOUVERNANCE CENTRALE & SÉCURITÉ FINTECH
    # ─────────────────────────────────────────────────────────────
    h9 = doc.add_paragraph()
    h9.paragraph_format.space_before = Pt(16)
    h9.paragraph_format.space_after = Pt(6)
    r9 = h9.add_run("9. Gouvernance Centrale, Portail Super-Admin (edoctor_admin) & Sécurité FinTech")
    r9.font.name = 'Arial'
    r9.font.size = Pt(14)
    r9.font.bold = True
    r9.font.color.rgb = DARK_NAVY

    p9 = doc.add_paragraph()
    p9.add_run(
        "Pour assurer une conformité réglementaire absolue et sécuriser les millions de flux financiers transitant par l'écosystème :\n\n"
        "1. Portail Super-Admin National (edoctor_admin) :\n"
        "Déployé sur https://admin-edoctor.ewaregroup.org (conteneurisé sous Docker / Nginx SPA et orchestré par Dokploy), "
        "cet outil de pilotage permet à la direction générale d'homologuer les hôpitaux et officines (validation des arrêtés MSHP et licences ONPT), "
        "d'auditer en temps réel le volume des téléconsultations et d'arbitrer les litiges usagers.\n\n"
        "2. Double Authentification (2FA OTP) & Passerelle Brevo :\n"
        "L'accès à toutes les interfaces (patients, praticiens, officines, super-administrateurs) est conditionné par un code temporaire à 6 chiffres "
        "expédié instantanément via la passerelle transactionnelle certifiée Brevo REST API v3 (contacts@ewaregroup.org).\n\n"
        "3. Protection Anti-Force Brute (Account Lockout) :\n"
        "Après 10 échecs de saisie de mot de passe consécutifs, le compte ciblé est automatiquement verrouillé pendant 30 minutes et une alerte "
        "de sécurité est expédiée au titulaire du compte.\n\n"
        "4. Sécurisation FinTech FedaPay en Production :\n"
        "Les paiements simulés sont formellement prohibés en production (403 Forbidden). Chaque notification webhook émise par FedaPay est "
        "validée par signature cryptographique HMAC (FEDAPAY_WEBHOOK_SECRET) avec un rate limiting de 60 requêtes/minute."
    )

    # ─────────────────────────────────────────────────────────────
    # CHAPITRE 10 : MATRICE RÉCAPITULATIVE DES RISQUES
    # ─────────────────────────────────────────────────────────────
    h10 = doc.add_paragraph()
    h10.paragraph_format.space_before = Pt(16)
    h10.paragraph_format.space_after = Pt(6)
    r10 = h10.add_run("10. Matrice Récapitulative des Risques & Garanties Opérationnelles")
    r10.font.name = 'Arial'
    r10.font.size = Pt(14)
    r10.font.bold = True
    r10.font.color.rgb = DARK_NAVY

    table_risks = doc.add_table(rows=1, cols=3)
    table_risks.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_risks.autofit = False

    risk_headers = ["Risque Analysé", "Contre-Mesure Technique eDoctor", "Garantie Partenaire"]
    r_hdr_cells = table_risks.rows[0].cells
    for i, title in enumerate(risk_headers):
        r_hdr_cells[i].text = title
        set_cell_background(r_hdr_cells[i], "132A45")
        set_cell_margins(r_hdr_cells[i], top=120, bottom=120, left=100, right=100)
        p = r_hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = 'Arial'
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    risk_rows = [
        ("Fraude de simulation de paiement", "Blocage strict en mode production + contrôle FedaPay out-of-band", "100% des fonds encaissés sont réels"),
        ("Attaque par injection ou altération Webhook", "Signature cryptographique HMAC (FEDAPAY_WEBHOOK_SECRET) + Rate Limit 60 req/min", "Traçabilité financière inviolable"),
        ("Piratage de compte soignant ou administrateur", "Double Authentification 2FA obligatoire par email (Brevo) + blocage après 10 échecs", "Protection maximale du secret médical"),
        ("Contournement de la plateforme médicale", "Ordonnance certifiée infalsifiable avec QR Code + anonymisation des échanges", "Fidélisation durable des soignants et patients"),
        ("Rupture de la chaîne du froid des médicaments", "Sacs isothermes scellés et acheminement rapide sous 45 minutes", "Qualité thérapeutique irréprochable"),
        ("Litiges financiers ou contestation de soins", "Module de litiges administré depuis le portail national Super-Admin", "Médiation rapide et transparence totale")
    ]

    for r_idx, r_val in enumerate(risk_rows):
        row = table_risks.add_row()
        bg_col = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, cell_txt in enumerate(r_val):
            c = row.cells[c_idx]
            c.text = cell_txt
            set_cell_background(c, bg_col)
            set_cell_margins(c, top=100, bottom=100, left=100, right=100)
            p = c.paragraphs[0]
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = CRIMSON_RED
            elif c_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.color.rgb = DARK_GRAY
            elif c_idx == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = PRIMARY_GREEN
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(9)

    doc.save(output_path)
    print(f"Master Dossier généré avec succès : {output_path}")

if __name__ == '__main__':
    target = os.path.join(os.getcwd(), 'DOSSIER_STRATEGIQUE_ET_PROCEDURES_OPERATIONNELLES_EDOCTOR.docx')
    create_master_dossier_docx(target)
