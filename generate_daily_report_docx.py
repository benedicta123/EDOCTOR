import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PRIMARY_GREEN = RGBColor(82, 153, 39)     # #529927
DARK_NAVY = RGBColor(19, 42, 69)          # #132A45
TEXT_MUTED = RGBColor(100, 116, 139)      # #64748B
DARK_GRAY = RGBColor(30, 41, 59)          # #1E293B

def set_cell_background(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_daily_report_docx(output_path):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

        hp = section.header.paragraphs[0]
        hp.text = "eDoctor (Togo Health Direct) — Rapport Journalier d'Activité"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if hp.runs:
            hp.runs[0].font.size = Pt(8.5)
            hp.runs[0].font.color.rgb = TEXT_MUTED

        fp = section.footer.paragraphs[0]
        fp.text = "Rapport du Soir • 30 septembre 2026 • eDoctor Platform"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if fp.runs:
            fp.runs[0].font.size = Pt(8.5)
            fp.runs[0].font.color.rgb = TEXT_MUTED

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("📋 Rapport Journalier d'Activité — 30 Septembre 2026")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = DARK_NAVY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("Bilan d'exécution du soir (conforme aux objectifs annoncés le matin)")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = TEXT_MUTED

    # Synthèse en 3 phrases
    p_box = doc.add_paragraph()
    p_box.paragraph_format.space_after = Pt(14)
    r_box = p_box.add_run("Rapport du soir (en 3 phrases) :\n\n")
    r_box.font.name = 'Arial'
    r_box.font.bold = True
    r_box.font.size = Pt(11)
    r_box.font.color.rgb = PRIMARY_GREEN

    r1 = p_box.add_run(
        "1. Aujourd'hui, nous avons mis en place avec succès le système d'alerte d'appel entrant en temps réel sur l'application Praticien (edoctor_praticient).\n\n"
        "2. Dès qu'un patient clique sur « Consulter maintenant », une modale d'alerte visuelle et sonore s'affiche désormais instantanément sur l'écran du médecin pour lui signaler la demande entrante.\n\n"
        "3. Le praticien dispose des actions immédiates pour « Accepter la consultation » et lancer directement la vidéo, ou « Refuser / Reporter » afin de garantir une prise en charge fluide, réactive et sécurisée."
    )
    r1.font.name = 'Arial'
    r1.font.size = Pt(10.5)
    r1.font.color.rgb = DARK_GRAY

    doc.save(output_path)
    print(f"Rapport journalier mis à jour : {output_path}")

if __name__ == '__main__':
    target = os.path.join(os.getcwd(), 'RAPPORT_JOURNALIER_2026_09_30.docx')
    generate_daily_report_docx(target)
