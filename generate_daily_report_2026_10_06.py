import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PRIMARY_GREEN = RGBColor(82, 153, 39)     # #529927
DARK_NAVY = RGBColor(19, 42, 69)          # #132A45
TEXT_MUTED = RGBColor(100, 116, 139)      # #64748B
DARK_GRAY = RGBColor(30, 41, 59)          # #1E293B

def set_cell_background(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_report():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

        hp = section.header.paragraphs[0]
        hp.text = "eDoctor (Togo Health Direct) — Rapport Journalier d'Activité"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if hp.runs:
            hp.runs[0].font.size = Pt(8.5)
            hp.runs[0].font.color.rgb = TEXT_MUTED

        fp = section.footer.paragraphs[0]
        fp.text = "Rapport du Soir • 06 octobre 2026 • eDoctor Platform"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if fp.runs:
            fp.runs[0].font.size = Pt(8.5)
            fp.runs[0].font.color.rgb = TEXT_MUTED

    # Titre
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("📋 Rapport Journalier d'Activité — 06 Octobre 2026")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(17)
    run_title.font.bold = True
    run_title.font.color.rgb = DARK_NAVY

    # Sous-titre
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("Bilan d'exécution du soir : Passerelle FedaPay, Commandes Pharmacies & Déploiement Production")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = TEXT_MUTED

    # Encadré : Synthèse en 3 phrases
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.8)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    p_box = cell.paragraphs[0]
    p_box.paragraph_format.space_after = Pt(6)
    r_head = p_box.add_run("🎯 Synthèse du soir (en 3 phrases) :\n\n")
    r_head.font.name = 'Arial'
    r_head.font.bold = True
    r_head.font.size = Pt(11)
    r_head.font.color.rgb = PRIMARY_GREEN

    r_text = p_box.add_run(
        "1. Aujourd'hui, nous avons intégré avec succès la passerelle de paiement Mobile Money FedaPay (T-Money Togo & Moov Money) avec gestion des sandbox de test et sécurisation des webhooks.\n\n"
        "2. Le module de commande en pharmacie a été fiabilisé avec le filtrage automatique des stocks à 100%, le calcul automatisé de la distance de livraison par GPS (formule Haversine) et l'harmonisation du statut des ordonnances déjà honorées.\n\n"
        "3. L'ensemble des services backend, migrations de base de données et routes d'API ont été déployés et validés en production sur l'infrastructure Dokploy du serveur VPS."
    )
    r_text.font.name = 'Arial'
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = DARK_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 1 : Détail des Réalisations
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    rh1 = h1.add_run("1. Réalisations Techniques Majeures")
    rh1.font.name = 'Arial'
    rh1.font.size = Pt(13)
    rh1.font.bold = True
    rh1.font.color.rgb = DARK_NAVY

    items = [
        ("Intégration Passerelle FedaPay & Téléconsultations :", 
         "Création du service FedaPayService et du contrôleur FedaPayPaymentController, migration de la base de données (order_id nullable dans payments), gestion des transactions sandbox avec webhooks et confirmation sécurisée."),
        ("Fiabilisation Stocks & Commandes Pharmacie :",
         "Correction de l'incohérence des prix à 0 FCFA (indisponibilités partielles), filtrage automatique strict des officines disposant de l'intégralité des médicaments prescrits en stock."),
        ("Sécurisation Calcul Distance de Livraison :",
         "Suppression du slider manuel modifiable par le patient (prévention anti-fraude) et mise en place du calcul de distance kilométrique automatique basé sur les coordonnées GPS officielles des quartiers de Lomé."),
        ("Ergonomie et Expérience Utilisateur (Mobile Patient) :",
         "Ajout d'une fenêtre de confirmation post-paiement avec lien direct « Voir commande », affichage de l'heure et date exacte sur chaque commande, et mise à jour du statut des ordonnances traitées en « Valide » (au lieu de « Valide / À traiter »)."),
        ("Mise en Production (Dokploy & VPS) :",
         "Synchronisation sur GitHub (origin main), reconstruction conteneurisée sur Dokploy Swarm, application des migrations et validation de l'état actif des routes d'API en production.")
    ]

    for title, desc in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        rt = p.add_run(f"• {title} ")
        rt.font.name = 'Arial'
        rt.font.bold = True
        rt.font.size = Pt(10)
        rt.font.color.rgb = PRIMARY_GREEN
        rd = p.add_run(desc)
        rd.font.name = 'Arial'
        rd.font.size = Pt(10)
        rd.font.color.rgb = DARK_GRAY

    # Section 2 : Tableau des Livrables et Statuts
    h2 = doc.add_heading(level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    rh2 = h2.add_run("2. Statut des Modules & Composants")
    rh2.font.name = 'Arial'
    rh2.font.size = Pt(13)
    rh2.font.bold = True
    rh2.font.color.rgb = DARK_NAVY

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(2.5), Inches(3.1), Inches(1.2)]

    hdr_cells = table.rows[0].cells
    headers = ["Composant / Écran", "Description / Évolution", "Statut"]
    for i, h_text in enumerate(headers):
        hdr_cells[i].width = widths[i]
        set_cell_background(hdr_cells[i], "132A45")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        r = p.add_run(h_text)
        r.font.name = 'Arial'
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    rows_data = [
        ("API Backend (Laravel 11)", "Passerelle FedaPay, webhook, eager loading order sur ordonnances", "Validé (Prod)"),
        ("Base de données (PostgreSQL)", "Migration payments (order_id nullable) appliquée", "Migré (Prod)"),
        ("App Mobile (Prescriptions)", "Calcul GPS automatique, badge 'Valide', modale post-achat", "Opérationnel"),
        ("App Mobile (Commandes)", "Affichage de la date et heure exacte (ex. 06/10 à 17h05)", "Opérationnel"),
        ("Infra Serveur (Dokploy/VPS)", "Conteneur recompilé, routes fedapay déclarées et actives", "Déployé"),
    ]

    for comp, desc, st in rows_data:
        row_cells = table.add_row().cells
        for i, val in enumerate([comp, desc, st]):
            row_cells[i].width = widths[i]
            set_cell_background(row_cells[i], "FFFFFF")
            set_cell_margins(row_cells[i], top=80, bottom=80, left=120, right=120)
            p = row_cells[i].paragraphs[0]
            r = p.add_run(val)
            r.font.name = 'Arial'
            r.font.size = Pt(9)
            if i == 2:
                r.font.bold = True
                r.font.color.rgb = PRIMARY_GREEN
            else:
                r.font.color.rgb = DARK_GRAY

    # Section 3 : Prochaines Étapes
    h3 = doc.add_heading(level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)
    rh3 = h3.add_run("3. Prochaines Priorités (Journée Suivante)")
    rh3.font.name = 'Arial'
    rh3.font.size = Pt(13)
    rh3.font.bold = True
    rh3.font.color.rgb = DARK_NAVY

    next_items = [
        "Passage en mode Live de FedaPay avec les clés de production pour les tests réels T-Money / Moov Money Togo.",
        "Génération des nouveaux APKs de démonstration pour les tests terrains utilisateurs.",
        "Surveillance des webhooks et du journal des transactions sur le tableau de bord Dokploy / FedaPay."
    ]

    for it in next_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        rt = p.add_run("→ ")
        rt.font.bold = True
        rt.font.color.rgb = PRIMARY_GREEN
        rd = p.add_run(it)
        rd.font.name = 'Arial'
        rd.font.size = Pt(10)
        rd.font.color.rgb = DARK_GRAY

    output_path = os.path.join(os.getcwd(), 'RAPPORT_JOURNALIER_2026_10_06.docx')
    doc.save(output_path)
    print(f"Rapport généré avec succès : {output_path}")

if __name__ == '__main__':
    generate_report()
